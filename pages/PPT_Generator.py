import streamlit as st
from groq import Groq
import os, json, io
from dotenv import load_dotenv
from pptx import Presentation
from shared import voice_or_text

load_dotenv()

st.set_page_config(page_title="PPT Generator", page_icon="🎞️", layout="wide")

st.markdown("""
<style>
    .block-container {
        padding-top: 2.5rem;
        max-width: 900px;
    }
    [data-testid="stChatMessage"] {
        border-radius: 18px;
        padding: 18px 22px;
        margin-bottom: 18px;
        border: 1px solid #2E2E2C;
        animation: fadeIn 0.3s ease-in;
    }
    @keyframes fadeIn {
        from { opacity: 0; transform: translateY(8px); }
        to { opacity: 1; transform: translateY(0); }
    }
    [data-testid="stChatInput"] textarea {
        border-radius: 14px;
        font-size: 15px;
    }
    .stButton button {
        border-radius: 10px;
        border: 1px solid #D97757;
        font-weight: 600;
        padding: 6px 18px;
        transition: all 0.2s ease;
    }
    .stButton button:hover {
        background-color: #D97757;
        border-color: #D97757;
        transform: scale(1.02);
    }
    [data-testid="stSidebar"] {
        background-color: #171716;
    }
    [data-testid="stSidebarNav"] a {
        border-radius: 8px;
        margin: 3px 8px;
        padding: 8px 10px;
        font-size: 14.5px;
    }
    [data-testid="stSidebarNav"] a:hover {
        background-color: #2A2A28;
    }
    [data-testid="stSidebarNav"] a[aria-current="page"] {
        background-color: #2E2420;
        border-left: 3px solid #D97757;
    }
    h1 {
        letter-spacing: -0.5px;
    }
</style>
""", unsafe_allow_html=True)
st.title("🎞️ PPT Generator")
st.caption("Turn any topic into a ready-to-present slide deck")

groq_client = Groq(api_key=os.getenv("GROQ_API_KEY"))

topic = voice_or_text("Enter your topic (e.g. 'RAG concept for my seminar'):", key="ppt")
num_slides = st.slider("Number of slides", min_value=3, max_value=15, value=8)

if st.button("Generate PPT"):
    if not topic:
        st.warning("Please enter a topic first.")
    else:
        with st.spinner("Generating slide content..."):
            prompt = f"""Create content for a {num_slides}-slide presentation on: "{topic}"

Return ONLY valid JSON (no markdown, no code fences, no explanation) in this exact format:
[
  {{"title": "Slide Title", "bullets": ["point 1", "point 2", "point 3"]}}
]

Make the first slide a title slide (one short bullet as subtitle).
Keep each bullet under 15 words. Make content accurate and well-organized."""

            response = groq_client.chat.completions.create(
                model="qwen/qwen3.8-27b",
                messages=[
                    {"role": "system", "content": "You are a presentation content generator. Respond with ONLY valid JSON, nothing else."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=900
            )

            raw = response.choices[0].message.content.strip()
            if raw.startswith("```"):
                raw = raw.split("```")[1]
                if raw.startswith("json"):
                    raw = raw[4:]

            try:
                slides_data = json.loads(raw)
            except json.JSONDecodeError:
                st.error("Couldn't parse the generated content. Please click Generate again.")
                st.stop()

        with st.spinner("Building PPT file..."):
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN
            from pptx.enum.shapes import MSO_SHAPE

            CHARCOAL = RGBColor(0x1A, 0x1A, 0x18)
            ORANGE = RGBColor(0xD9, 0x77, 0x57)
            INK = RGBColor(0x26, 0x26, 0x24)
            MUTED = RGBColor(0x6B, 0x6B, 0x66)
            WHITE = RGBColor(0xFF, 0xFF, 0xFF)

            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            blank_layout = prs.slide_layouts[6]

            for idx, item in enumerate(slides_data):
                slide = prs.slides.add_slide(blank_layout)
                is_title = (idx == 0)

                bg = slide.background
                bg.fill.solid()
                bg.fill.fore_color.rgb = CHARCOAL if is_title else WHITE

                if is_title:
                    box = slide.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11.33), Inches(1.5))
                    tf = box.text_frame
                    tf.word_wrap = True
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = item["title"]
                    run.font.size = Pt(40)
                    run.font.bold = True
                    run.font.color.rgb = WHITE
                    run.font.name = "Cambria"

                    bullets = item.get("bullets", [])
                    if bullets:
                        sbox = slide.shapes.add_textbox(Inches(1.5), Inches(4.15), Inches(10.33), Inches(0.8))
                        stf = sbox.text_frame
                        stf.word_wrap = True
                        sp = stf.paragraphs[0]
                        sp.alignment = PP_ALIGN.CENTER
                        srun = sp.add_run()
                        srun.text = bullets[0]
                        srun.font.size = Pt(18)
                        srun.font.italic = True
                        srun.font.color.rgb = ORANGE
                        srun.font.name = "Calibri"
                else:
                    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.14),
                                                 prs.slide_height)
                    bar.fill.solid()
                    bar.fill.fore_color.rgb = ORANGE
                    bar.line.fill.background()

                    tbox = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.9), Inches(1.0))
                    tf = tbox.text_frame
                    tf.word_wrap = True
                    p = tf.paragraphs[0]
                    run = p.add_run()
                    run.text = item["title"]
                    run.font.size = Pt(28)
                    run.font.bold = True
                    run.font.color.rgb = INK
                    run.font.name = "Cambria"

                    bbox = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.3))
                    btf = bbox.text_frame
                    btf.word_wrap = True
                    for i, bullet in enumerate(item.get("bullets", [])):
                        p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
                        p.space_after = Pt(14)
                        run = p.add_run()
                        run.text = "•  " + bullet
                        run.font.size = Pt(17)
                        run.font.color.rgb = MUTED
                        run.font.name = "Calibri"

            buffer = io.BytesIO()
            prs.save(buffer)
            buffer.seek(0)

        st.success(f"✅ Generated {len(slides_data)} slides on '{topic}'")

        st.download_button(
            label="📥 Download PPT",
            data=buffer,
            file_name=f"{topic.replace(' ', '_')}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )

        st.markdown("### Preview")
        for item in slides_data:
            with st.expander(item["title"]):
                for b in item.get("bullets", []):
                    st.markdown(f"- {b}")