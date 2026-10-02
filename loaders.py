from pypdf import PdfReader
from pptx import Presentation
from docx import Document


def load_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        text += page.extract_text() + "\n"
    return text


def load_pptx(file_path):
    prs = Presentation(file_path)
    text = ""
    for slide_num, slide in enumerate(prs.slides, 1):
        text += f"\n--- Slide {slide_num} ---\n"
        for shape in slide.shapes:
            if shape.has_text_frame:
                text += shape.text_frame.text + "\n"
    return text


def load_docx(file_path):
    doc = Document(file_path)
    text = "\n".join([para.text for para in doc.paragraphs])
    return text


def load_txt(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def load_file(file_path):
    """File extension பாத்து correct loader-ஐ call பண்ணும்"""
    ext = file_path.lower().split(".")[-1]

    loaders = {
        "pdf": load_pdf,
        "pptx": load_pptx,
        "docx": load_docx,
        "txt": load_txt,
    }

    if ext not in loaders:
        raise ValueError(f"Unsupported file type: {ext}")

    return loaders[ext](file_path)